"""Offline ResearchSnapshot to XLSX export using openpyxl."""
from __future__ import annotations
import json, os, re, uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

class ExcelExportError(RuntimeError): pass
EXPECTED_SHEETS = ["Overview","Market_Daily","SEC_Filings","SEC_Facts","Sources","Data_Quality","Run_Record"]
_FORMULA = re.compile(r"^[=+\-@]")
_URL = re.compile(r"^https?://[^\s]+$", re.I)
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:?\d{2})?)?$")
_SYMBOL = re.compile(r"^[A-Z][A-Z0-9.\-]{0,9}$")

def _validate_snapshot(s: dict[str, Any]) -> None:
    required={"schema_version","snapshot_id","as_of","security","market_data","sec_facts","sec_filings","sources","data_quality","run_record"}
    missing=sorted(required-set(s))
    if missing: raise ExcelExportError(f"snapshot missing required fields: {', '.join(missing)}")
    if s.get("schema_version")!="1.1": raise ExcelExportError("schema_version must be 1.1; schema 1.0 is not overwritten")
    if s.get("snapshot_type")!="single_stock_research": raise ExcelExportError("snapshot_type must be single_stock_research")
    if s.get("data_mode") not in {"synthetic","saved_snapshot"}: raise ExcelExportError("data_mode must be synthetic or saved_snapshot")
    for k in ("snapshot_id","as_of"):
        if not isinstance(s.get(k),str) or not s[k].strip(): raise ExcelExportError(f"{k} must be non-empty text")
    if not isinstance(s["security"],dict) or not isinstance(s["security"].get("symbol"),str) or not s["security"]["symbol"].strip(): raise ExcelExportError("snapshot.security.symbol is required text")
    if not _SYMBOL.fullmatch(s["security"]["symbol"].strip().upper()): raise ExcelExportError("snapshot.security.symbol is invalid")
    def valid_iso(value: Any) -> bool:
        if not isinstance(value, str) or not _DATE.match(value): return False
        try:
            if "T" not in value and " " not in value: date.fromisoformat(value)
            else: datetime.fromisoformat(value.replace("Z", "+00:00"))
            return True
        except ValueError: return False
    if not valid_iso(s["as_of"]): raise ExcelExportError("as_of must be a valid ISO date or datetime")
    if s.get("created_at") is not None and not valid_iso(s["created_at"]): raise ExcelExportError("created_at must be a valid ISO date or datetime")
    if not isinstance(s["market_data"],dict): raise ExcelExportError("market_data must be an object")
    if not isinstance(s["market_data"].get("bars"), list) or any(not isinstance(r, dict) for r in s["market_data"]["bars"]): raise ExcelExportError("market_data.bars must be a list of objects")
    for k in ("sec_facts","sec_filings","sources"):
        if not isinstance(s[k],list) or any(not isinstance(r,dict) for r in s[k]): raise ExcelExportError(f"{k} must be a list of objects")
    if not isinstance(s["data_quality"],dict) or not isinstance(s["run_record"],dict): raise ExcelExportError("data_quality and run_record must be objects")
    if type(s["run_record"].get("network_executed")) is not bool or s["run_record"]["network_executed"] is not False: raise ExcelExportError("export requires run_record.network_executed to be boolean false")

def _local_root(p: Path) -> Path:
    p=p.absolute()
    for parent in (p,*p.parents):
        if parent.name==".local_data": return parent
    raise ExcelExportError("path must be inside a .local_data directory")

def _safe_local_path(p: Path, root: Path) -> None:
    root=root.absolute()
    if os.path.lexists(root) and root.is_symlink(): raise ExcelExportError(f".local_data root symlink is not allowed: {root}")
    if os.path.lexists(p) and p.is_symlink(): raise ExcelExportError(f"symlink output path is not allowed: {p}")
    try: rel=p.absolute().relative_to(root)
    except ValueError as e: raise ExcelExportError("path must be inside .local_data") from e
    cur=root
    for part in rel.parts[:-1]:
        cur/=part
        if os.path.lexists(cur) and cur.is_symlink(): raise ExcelExportError(f"symlink path component is not allowed: {cur}")

def _value(v: Any) -> Any:
    if v is None or isinstance(v,(bool,int,float,date,datetime)): return v
    if isinstance(v,(list,dict)): return json.dumps(v,ensure_ascii=False,sort_keys=True)
    t=str(v); return "'"+t if _FORMULA.match(t) else t

def _dated(v: Any) -> Any:
    if not isinstance(v,str) or not _DATE.match(v): return _value(v)
    try:
        if "T" not in v and " " not in v: return date.fromisoformat(v)
        # Excel stores naive datetimes; preserve the displayed UTC instant while
        # removing tzinfo so openpyxl can serialize it safely.
        parsed = datetime.fromisoformat(v.replace("Z","+00:00"))
        return parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed
    except ValueError: return _value(v)

def _table(ws, title, headers, rows, s):
    ws.sheet_view.showGridLines=False; ws["A1"]=_value(title); ws["A1"].font=Font(name="Arial",size=14,bold=True,color="17365D")
    ws["A2"]=_value(f"Snapshot {s['snapshot_id']} | as_of {s['as_of']} | mode {s.get('data_mode','unknown')}"); ws["A2"].font=Font(name="Arial",size=9,italic=True,color="5B6573")
    fill=PatternFill("solid",fgColor="17365D"); side=Side(style="thin",color="D9E2F3")
    for c,h in enumerate(headers,1):
        x=ws.cell(4,c,_value(h)); x.fill=fill; x.font=Font(name="Arial",size=10,bold=True,color="FFFFFF"); x.alignment=Alignment(horizontal="center")
    for r,row in enumerate(rows,5):
        for c,v in enumerate(row,1):
            x=ws.cell(r,c,_dated(v) if isinstance(v,str) else _value(v))
            if isinstance(x.value,datetime): x.number_format="yyyy-mm-dd hh:mm:ss"
            elif isinstance(x.value,date): x.number_format="yyyy-mm-dd"
            elif isinstance(x.value,(int,float)) and not isinstance(x.value,bool): x.number_format="#,##0.########"
            if isinstance(v,str) and _URL.match(v): x.hyperlink=v; x.style="Hyperlink"
    mr=max(4,len(rows)+4); mc=max(1,len(headers))
    for row in ws.iter_rows(min_row=4,max_row=mr,min_col=1,max_col=mc):
        for x in row: x.border=Border(bottom=side); x.alignment=Alignment(vertical="center",wrap_text=x.row>4)
    ws.freeze_panes="A5"
    ws.auto_filter.ref=f"A4:{get_column_letter(mc)}{max(4, len(rows)+4)}"
    for c in range(1,mc+1): ws.column_dimensions[get_column_letter(c)].width=min(42,max(12,max(len(str(ws.cell(r,c).value or "")) for r in range(1,mr+1))+2))

def _tables(s):
    sec,md,dq,rr=s["security"],s["market_data"],s["data_quality"],s["run_record"]; summary=s.get("summary",{}); latest=md.get("latest_bar") or {}
    overview=[["Symbol",sec.get("symbol"),"Company",sec.get("company_name")],["CIK",sec.get("cik"),"Exchange",sec.get("exchange")],["Currency",sec.get("currency"),"Schema version",s.get("schema_version")],["Snapshot ID",s.get("snapshot_id"),"As Of",s.get("as_of")],["Created At",s.get("created_at"),"Data Mode",s.get("data_mode")],["Latest market observation",summary.get("latest_market_observation"),"Market source",md.get("source")],["Latest close",latest.get("close"),"Latest volume",latest.get("volume")],["Market Bar Count",len(md.get("bars",[])),"Filing Count",len(s.get("sec_filings",[]))],["Fact Count",len(s.get("sec_facts",[])),"Source Count",len(s.get("sources",[]))],["Data quality",dq.get("status"),"Evidence status",summary.get("evidence_status")],["Network Executed",rr.get("network_executed"),"Model Calls",rr.get("model_calls")],["Disclaimer", "For research assistance only. Not investment advice.", "Delivery", "Generated report does not mean an email was sent."],["Export Mode", "Saved Snapshot / Offline Export", "Realtime", "Not real-time"]]
    quality=[]
    for k in ("status","missing_fields","missing_sources","warnings","conflicts","point_in_time_filtered_count","total_excluded_count","filter_counts"):
        v=dq.get(k)
        if isinstance(v,list): quality.extend([[k,x] for x in v] or [[k,None]])
        elif isinstance(v,dict): quality.extend([[f"{k}.{a}",b] for a,b in v.items()] or [[k,None]])
        else: quality.append([k,v])
    run=[["Run ID",rr.get("run_id")],["Status",rr.get("status")],["Started At",rr.get("started_at")],["Completed At",rr.get("completed_at")],["Input Mode",rr.get("input_mode")],["Network Executed",rr.get("network_executed")],["Model Calls",rr.get("model_calls")],["File Hash Verification",rr.get("file_hash_verification")],["Identity Verification",rr.get("identity_verification")],["Snapshot Write",rr.get("snapshot_write")],["Point-in-Time Filter",rr.get("point_in_time_filter")],["Selected Files",rr.get("selected_files")],["Steps",rr.get("steps")],["Warnings",rr.get("warnings")],["Errors",rr.get("errors")],["Snapshot ID",s.get("snapshot_id")],["Schema version",s.get("schema_version")],["Created at",s.get("created_at")],["As of",s.get("as_of")],["Data mode",s.get("data_mode")],["Data quality status",dq.get("status")]]
    known_run_keys={"run_id","status","started_at","completed_at","input_mode","network_executed","model_calls","file_hash_verification","identity_verification","snapshot_write","point_in_time_filter","selected_files","steps","warnings","errors"}
    run.extend([[k,v] for k,v in rr.items() if k not in known_run_keys])
    return {
      "Overview":(["Field","Value","Field","Value"],overview),
      "Market_Daily":(["Symbol","Date","Timestamp","Open","High","Low","Close","Adjusted Close","Volume","Currency","Source","Retrieved At","Adjustment"],[[r.get("symbol",sec.get("symbol")),r.get("as_of",r.get("date")),r.get("timestamp"),r.get("open"),r.get("high"),r.get("low"),r.get("close"),r.get("adjusted_close"),r.get("volume"),r.get("currency",md.get("currency")),r.get("source",md.get("source")),r.get("retrieved_at",md.get("retrieved_at")),md.get("adjustment")] for r in md.get("bars",[])]),
      "SEC_Filings":(["Form","Filing date","Report date","Acceptance datetime","Accession","Primary document","Published at","Source","Source URL","Retrieved at","CIK"],[[r.get("form"),r.get("filing_date"),r.get("report_date"),r.get("acceptance_datetime"),r.get("accession_number"),r.get("primary_document"),r.get("published_at"),r.get("source"),r.get("source_url",r.get("filing_url")),r.get("retrieved_at"),r.get("cik",sec.get("cik"))] for r in s["sec_filings"]]),
      "SEC_Facts":(["Taxonomy","Tag","Label","Value","Unit","Currency","Period start","Period end","Filed At","Available At","Form","Frame","Accession","Source","Source URL","Retrieved At"],[[r.get("taxonomy"),r.get("tag"),r.get("label"),r.get("value"),r.get("unit"),r.get("currency"),r.get("period_start"),r.get("period_end"),r.get("filed_at"),r.get("available_at"),r.get("form"),r.get("frame"),r.get("accession_number"),r.get("source"),r.get("source_url"),r.get("retrieved_at")] for r in s["sec_facts"]]),
      "Sources":(["Source ID","Type","Provider","Source URL","Published at","Data as of","Retrieved at","Raw SHA-256","Normalized SHA-256","Input file"],[[r.get("source_id"),r.get("source_type"),r.get("provider"),r.get("source_url"),r.get("published_at"),r.get("data_as_of"),r.get("retrieved_at"),r.get("raw_sha256"),r.get("normalized_sha256"),r.get("input_file")] for r in s["sources"]]),
      "Data_Quality":(["Field","Value"],quality),"Run_Record":(["Field","Value"],run)}

def export_snapshot_to_excel(snapshot: dict[str,Any], output_path: str|os.PathLike[str], *, overwrite: bool=False)->Path:
    _validate_snapshot(snapshot); dest=Path(output_path).expanduser(); root=_local_root(dest)
    if dest.suffix.lower()!=".xlsx": raise ExcelExportError("output must have .xlsx extension")
    _safe_local_path(dest,root)
    if os.path.lexists(dest) and not overwrite: raise ExcelExportError(f"output already exists (use overwrite=True): {dest}")
    dest.parent.mkdir(parents=True,exist_ok=True); tmp=dest.parent/f".{dest.name}.{uuid.uuid4().hex}.tmp.xlsx"
    titles={"Overview":"FinSight Research Snapshot","Market_Daily":"Saved market observations","SEC_Filings":"SEC filings available by cutoff","SEC_Facts":"SEC Company Facts available by cutoff","Sources":"Source and version metadata","Data_Quality":"Data quality and filtering","Run_Record":"Offline run and data-quality record"}
    try:
        wb=Workbook(); wb.remove(wb.active); tables=_tables(snapshot)
        for name in EXPECTED_SHEETS:
            ws=wb.create_sheet(name); headers,rows=tables[name]; _table(ws,titles[name],headers,rows,snapshot)
        wb.save(tmp); check=load_workbook(tmp,read_only=True,data_only=False)
        if check.sheetnames!=EXPECTED_SHEETS: raise ExcelExportError(f"reopen sheet verification failed: {check.sheetnames}")
        check.close(); os.replace(tmp,dest)
    except ExcelExportError: raise
    except Exception as exc: raise ExcelExportError(f"openpyxl export failed: {exc}") from exc
    finally:
        if tmp.exists(): tmp.unlink()
    return dest

def export_snapshot_file(snapshot_path: str|os.PathLike[str], output_path: str|os.PathLike[str], **kwargs:Any)->Path:
    path=Path(snapshot_path); _safe_local_path(path,_local_root(path))
    if path.is_symlink(): raise ExcelExportError("snapshot symlink is not allowed")
    try: snapshot=json.loads(path.read_text(encoding="utf-8"))
    except (OSError,json.JSONDecodeError) as exc: raise ExcelExportError(f"cannot read snapshot JSON: {path}") from exc
    if not isinstance(snapshot,dict): raise ExcelExportError("snapshot JSON must contain an object")
    return export_snapshot_to_excel(snapshot,output_path,**kwargs)
