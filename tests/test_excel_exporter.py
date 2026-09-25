import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from finsight_research.excel_exporter import ExcelExportError, export_snapshot_file


ROOT = Path(__file__).resolve().parents[1]


class ExcelExporterTests(unittest.TestCase):
    def _snapshot(self, temp: Path) -> Path:
        temp = temp / ".local_data"
        temp.mkdir()
        snapshot = {
            "schema_version": "1.1", "snapshot_id": "snapshot-fixture", "snapshot_type": "single_stock_research",
            "created_at": "2024-01-04T00:00:00Z", "as_of": "2024-01-03T00:00:00Z", "data_mode": "synthetic",
            "security": {"symbol": "TEST", "company_name": "Synthetic Example Corp", "currency": "USD"},
            "market_data": {"source": "fixture", "bars": [{"symbol": "TEST", "as_of": "2024-01-02", "close": 12.5, "volume": 1000}]},
            "sec_facts": [{"tag": "Revenue", "value": 10, "unit": "USD", "available_at": "2024-01-02T00:00:00Z", "source": "fixture"}],
            "sec_filings": [{"form": "10-K", "filing_date": "2024-01-02", "published_at": "2024-01-02", "source": "fixture"}],
            "summary": {"evidence_status": "complete", "available_fact_count": 1, "available_filing_count": 1},
            "sources": [{"source_id": "fixture", "provider": "fixture", "retrieved_at": "2024-01-04T00:00:00Z"}],
            "data_quality": {"status": "completed", "missing_sources": [], "filter_counts": {}},
            "run_record": {"network_executed": False, "model_calls": 0, "status": "completed"},
        }
        path = temp / "snapshot.json"
        path.write_text(json.dumps(snapshot), encoding="utf-8")
        return path

    def test_exports_xlsx_with_expected_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / ".local_data" / "research.xlsx"
            export_snapshot_file(self._snapshot(root), output)
            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 0)
            self.assertEqual(output.read_bytes()[:2], b"PK")

    def test_rejects_network_snapshot(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self._snapshot(root)
            data = json.loads(path.read_text())
            for value in (True, 1, "true", None):
                data["run_record"]["network_executed"] = value
                path.write_text(json.dumps(data))
                with self.assertRaises(ExcelExportError):
                    export_snapshot_file(path, root / ".local_data" / "research.xlsx")

    def test_requires_explicit_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self._snapshot(root)
            output = root / ".local_data" / "research.xlsx"
            export_snapshot_file(path, output)
            with self.assertRaises(ExcelExportError):
                export_snapshot_file(path, output)
            export_snapshot_file(path, output, overwrite=True)

    def test_rejects_invalid_snapshot_structure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self._snapshot(root)
            data = json.loads(path.read_text())
            for field, value in (("schema_version", "1.0"), ("snapshot_type", "other"), ("data_mode", "live"), ("as_of", ""), ("sec_facts", {})):
                altered = dict(data)
                altered[field] = value
                path.write_text(json.dumps(altered))
                with self.assertRaises(ExcelExportError):
                    export_snapshot_file(path, root / ".local_data" / "bad.xlsx")
            path.write_text(json.dumps(data))

    def test_rejects_symlink_output_and_path_escape(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self._snapshot(root)
            target = root / ".local_data" / "missing.xlsx"
            link = root / ".local_data" / "link.xlsx"
            link.symlink_to(target)
            with self.assertRaises(ExcelExportError):
                export_snapshot_file(path, link)
            with self.assertRaises(ExcelExportError):
                export_snapshot_file(path, root / "outside.xlsx")

    def test_escapes_formula_like_strings(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = self._snapshot(root)
            data = json.loads(path.read_text())
            data["sources"][0]["source_url"] = "=HYPERLINK(\"https://invalid\")"
            path.write_text(json.dumps(data))
            output = root / ".local_data" / "safe.xlsx"
            export_snapshot_file(path, output)
            with zipfile.ZipFile(output) as archive:
                xml = "".join(archive.read(name).decode("utf-8", "ignore") for name in archive.namelist() if name.startswith("xl/") and name.endswith(".xml"))
            self.assertNotIn("<x:f>", xml)
            self.assertIn("=HYPERLINK", xml)

    def test_workbook_contract_filters_and_field_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory); path=self._snapshot(root); out=root/".local_data"/"report.xlsx"; export_snapshot_file(path,out)
            from openpyxl import load_workbook
            wb=load_workbook(out, read_only=False, data_only=False)
            self.assertEqual(wb.sheetnames,["Overview","Market_Daily","SEC_Filings","SEC_Facts","Sources","Data_Quality","Run_Record"])
            for name in ("Market_Daily","SEC_Filings","SEC_Facts","Sources","Data_Quality","Run_Record"):
                self.assertIsNotNone(wb[name].auto_filter.ref); self.assertEqual(wb[name].freeze_panes,"A5")
            headers=[c.value for c in wb["SEC_Facts"][4]]
            self.assertIn("Filed At", headers); self.assertIn("Available At", headers)
            values=[c.value for row in wb["Overview"].iter_rows() for c in row]
            self.assertIn("For research assistance only. Not investment advice.", values)
            self.assertIn("Generated report does not mean an email was sent.", values)
            wb.close()

    def test_rejects_bad_as_of_symbol_and_bars(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory); path=self._snapshot(root); data=json.loads(path.read_text())
            for key,value in (("as_of","bad"),("created_at","bad"),("market_data",{"bars":{}}),("security",{"symbol":"bad symbol"})):
                altered=dict(data); altered[key]=value; path.write_text(json.dumps(altered))
                with self.assertRaises(ExcelExportError): export_snapshot_file(path,root/".local_data"/"bad.xlsx")

    def test_output_parent_symlink_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory); path=self._snapshot(root); outside=Path(directory)/"outside"; outside.mkdir(); (root/".local_data"/"linkdir").symlink_to(outside,target_is_directory=True)
            with self.assertRaises(ExcelExportError): export_snapshot_file(path,root/".local_data"/"linkdir"/"x.xlsx")
